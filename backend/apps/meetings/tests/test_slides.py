import shutil
import tempfile
from io import BytesIO
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from pptx import Presentation
from pypdf import PdfWriter

from apps.meetings.models import SharedSlides
from apps.meetings.services.slides import PPTX_MEDIA_TYPE


def _pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=100)
    writer.add_blank_page(width=200, height=100)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[0])
    first.notes_slide.notes_text_frame.text = "Welcome the room"
    second = presentation.slides.add_slide(presentation.slide_layouts[0])
    second.notes_slide.notes_text_frame.text = "Ask for questions"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


class SlidesApiTests(TestCase):
    def setUp(self):
        self.media = tempfile.mkdtemp()
        self.settings_override = override_settings(MEDIA_ROOT=self.media, MEETINGS_API_KEY="")
        self.settings_override.enable()
        lock_patch = patch("apps.meetings.services.slides.set_screen_share_locked")
        self.lock_screen = lock_patch.start()
        self.addCleanup(lock_patch.stop)

    def tearDown(self):
        self.settings_override.disable()
        shutil.rmtree(self.media, ignore_errors=True)

    def _upload(self, room="demo", identity="p_owner", name="Nico"):
        upload = SimpleUploadedFile("deck.pdf", _pdf_bytes(), content_type="application/pdf")
        return self.client.post(
            "/api/slides",
            {
                "room_name": room,
                "owner_identity": identity,
                "owner_name": name,
                "file": upload,
            },
        )

    def test_upload_page_and_stop(self):
        created = self._upload()
        self.assertEqual(created.status_code, 201)
        deck = created.json()["deck"]
        self.assertEqual(deck["pageCount"], 2)
        self.assertEqual(deck["page"], 1)
        self.assertEqual(deck["format"], "pdf")
        self.assertEqual(deck["ownerName"], "Nico")
        self.assertFalse(deck["hasNotes"])
        self.assertEqual(created.json()["notes"], [])

        current = self.client.get("/api/slides", {"roomName": "demo"})
        self.assertEqual(current.json()["deck"]["deckId"], deck["deckId"])
        self.assertNotIn("converting", current.json())

        pdf = self.client.get(f"/api/slides/{deck['deckId']}/file")
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        body = b"".join(pdf.streaming_content)
        self.assertTrue(body.startswith(b"%PDF-"))

        denied = self.client.post(
            f"/api/slides/{deck['deckId']}/page",
            data={"page": 2, "owner_identity": "p_other"},
            content_type="application/json",
        )
        self.assertEqual(denied.status_code, 403)

        moved = self.client.post(
            f"/api/slides/{deck['deckId']}/page",
            data={"page": 2, "owner_identity": "p_owner"},
            content_type="application/json",
        )
        self.assertEqual(moved.status_code, 200)
        self.assertEqual(moved.json()["deck"]["page"], 2)

        stopped = self.client.post(
            f"/api/slides/{deck['deckId']}/stop",
            data={"owner_identity": "p_owner"},
            content_type="application/json",
        )
        self.assertEqual(stopped.status_code, 200)
        self.assertIsNone(self.client.get("/api/slides", {"roomName": "demo"}).json()["deck"])
        self.assertFalse(SharedSlides.objects.get(id=deck["deckId"]).active)
        self.lock_screen.assert_any_call("demo", True)
        self.lock_screen.assert_any_call("demo", False)

    def test_second_upload_replaces_the_room_deck(self):
        first = self._upload().json()["deck"]
        second = self._upload().json()["deck"]
        current = self.client.get("/api/slides", {"roomName": "demo"}).json()["deck"]
        self.assertEqual(current["deckId"], second["deckId"])
        self.assertEqual(self.client.get(f"/api/slides/{first['deckId']}/file").status_code, 404)

    def test_rejects_non_pdf(self):
        upload = SimpleUploadedFile("notes.txt", b"hello", content_type="text/plain")
        response = self.client.post(
            "/api/slides",
            {"room_name": "demo", "owner_identity": "p_owner", "file": upload},
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["detail"], "Choose a PDF or PowerPoint file")

    def test_pptx_is_stored_and_keeps_speaker_notes_for_the_owner(self):
        upload = SimpleUploadedFile(
            "deck.pptx",
            _pptx_bytes(),
            content_type=PPTX_MEDIA_TYPE,
        )
        created = self.client.post(
            "/api/slides",
            {
                "room_name": "demo",
                "owner_identity": "p_owner",
                "owner_name": "Nico",
                "file": upload,
            },
        )
        self.assertEqual(created.status_code, 201)
        body = created.json()
        deck = body["deck"]
        self.assertEqual(deck["format"], "pptx")
        self.assertEqual(deck["pageCount"], 3)
        self.assertTrue(deck["hasNotes"])
        self.assertEqual(body["notes"], ["Welcome the room", "Ask for questions", ""])
        self.assertNotIn("notes", deck)

        stored = self.client.get(f"/api/slides/{deck['deckId']}/file")
        self.assertEqual(stored.status_code, 200)
        self.assertEqual(stored["Content-Type"], PPTX_MEDIA_TYPE)
        payload = b"".join(stored.streaming_content)
        self.assertTrue(payload.startswith(b"PK"))
        self.assertFalse(payload.startswith(b"%PDF-"))

        public = self.client.get("/api/slides", {"roomName": "demo"})
        self.assertNotIn("Welcome the room", public.content.decode())
        self.assertEqual(public.json()["deck"]["deckId"], deck["deckId"])
        self.assertTrue(public.json()["deck"]["hasNotes"])

        denied = self.client.get(
            f"/api/slides/{deck['deckId']}/notes",
            {"owner_identity": "p_other"},
        )
        self.assertEqual(denied.status_code, 403)
        allowed = self.client.get(
            f"/api/slides/{deck['deckId']}/notes",
            {"owner_identity": "p_owner"},
        )
        self.assertEqual(allowed.status_code, 200)
        self.assertEqual(allowed.json()["notes"], ["Welcome the room", "Ask for questions", ""])

    def test_rejects_invalid_pptx(self):
        upload = SimpleUploadedFile(
            "deck.pptx",
            b"not a deck",
            content_type=PPTX_MEDIA_TYPE,
        )
        response = self.client.post(
            "/api/slides",
            {"room_name": "demo", "owner_identity": "p_owner", "file": upload},
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["detail"], "Could not read that PowerPoint file")
        self.assertIsNone(self.client.get("/api/slides", {"roomName": "demo"}).json()["deck"])
